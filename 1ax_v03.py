import copy
import copy
import os
import matplotlib.pyplot as plt
import numpy as np
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch, cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PyPDF2 import PdfFileWriter, PdfFileReader
from reportlab.platypus import Frame, Image
from pandas import read_excel
from numpy import nan
import io
import fnmatch
from PyPDF2 import PdfFileWriter, PdfFileReader
from pandas import to_datetime
import openpyxl
from pandas import read_excel
from numpy import nan
# from reportlab.graphics.shapes import Rect,Line
# from matplotlib.transforms import Affine2D
# import matplotlib.ticker as plticker





def grafik_darstellung(versuch_list_excel,probe):
    import os
    from sklearn.metrics import r2_score
    from matplotlib.patches import Rectangle
    from matplotlib.transforms import Bbox
    roh_data_path = os.getcwd() + "\\roh_1ax"

    versuch_nummer = read_excel(roh_data_path+ "/" +str(probe)+".xls", sheet_name="Daten",header=None).round(decimals=6)


    durchmesser_cm = 11.40
    anfangshöhe_cm = 25.00
    anfangsvolumen_cm3 = round(np.pi*(durchmesser_cm*0.5)**2*anfangshöhe_cm,2)
    weg_intern = versuch_nummer[2]
    weg_ab_kontakt = weg_intern - weg_intern[0]
    kraft_kN = versuch_nummer[1]
    massgebliche_querschnitt_m2 = anfangsvolumen_cm3 / (anfangshöhe_cm - weg_ab_kontakt / 10) / 10000
    spannung_kn_m2 = kraft_kN / massgebliche_querschnitt_m2
    stauchung = ((weg_ab_kontakt / 10) / anfangshöhe_cm) * 100

    curve = np.polyfit(stauchung, spannung_kn_m2, 10)
    poly = np.poly1d(curve)
    new_y = [poly(i) for i in stauchung]
    r2 = r2_score(spannung_kn_m2,new_y)
    return_radius = spannung_kn_m2.to_list()
    return_radius_index = return_radius.index(max(spannung_kn_m2))
    return_stauchung = stauchung[return_radius_index]
    plt.rcParams["figure.figsize"] = (12, 12)

    fig = plt.figure()
    ax1 = fig.add_subplot(111)
    plt.rc('xtick', labelsize=12)
    plt.rc('ytick', labelsize=12)

    # plt.setp(ax1, xticks=range(0, int(max(stauchung)), 5), yticks=range(0,int(max(spannung_kn_m2)),5))
    # ax2 = fig.add_subplot(212)

    # fig, axes = plt.subplots(nrows=2, ncols=1)
    # ax1, ax2 = axes.flatten()
    ax1.plot(stauchung,spannung_kn_m2,marker = "s",markevery=15,linewidth=2.5,color = "red",markersize=7,label = "Messung")
    ax1.plot(stauchung, new_y, linewidth=2.5, color="b",label ="Polynomiale Regression" )
    ax1.plot([],[]," ",label = f"R\u00b2 : {round(r2,3)} " )
    ax1.set_xlabel("Stauchung \u03B5 [%]",fontsize=20,fontname="Arial")
    ax1.set_ylabel("Einaxiale Druckspannung "+ r"$\sigma$" + " [kN/m²]",fontsize=20,fontname="Arial")
    ax1.set_xlim(0,int(max(stauchung))+1)
    ax1.set_ylim(0,int(max(spannung_kn_m2))+5)
    ax1.xaxis.set_tick_params(labelsize=16)
    ax1.yaxis.set_tick_params(labelsize=16)

    # fig.patches.extend([plt.Rectangle((0.2,0.2), 0.1, 0.1,
    #                                   fill=True, color='g', alpha=0.5, zorder=1000,
    #                                   transform=fig.transFigure, figure=fig)])
    # x = 0.1  # Dikdörtgenin sol alt köşesinin x koordinatı (0-1 aralığında)
    # y = 0.6  # Dikdörtgenin sol alt köşesinin y koordinatı (0-1 aralığında)
    # width = 0.3  # Dikdörtgenin genişliği (0-1 aralığında)
    # height = 0.3  # Dikdörtgenin yüksekliği (0-1 aralığında)
    #
    # rectangle = Rectangle((x,y), width,height, edgecolor='red', facecolor='none')
    # ax1.add_patch(rectangle)
    # fig_width, fig_height = fig.get_size_inches()  # Grafik boyutlarını alır
    # bbox = Bbox.from_bounds(x * fig_width, y * fig_height, width * fig_width,
    #                         height * fig_height)  # Ölçeklenmiş dikdörtgen koordinatlarını oluşturur
    # rectangle.set_bounds(*bbox.bounds)  # Dikdörtgenin koordinatlarını günceller

    # ax2.plot([],[]," ",label = f"R\u00b2 : {round(r2,3)} " )

    # ax1 = plt.plot(stauchung, new_y, marker="s", markevery=15,linewidth=1,color = "b")
    # ax1 = plt.plot(stauchung, new_y, linewidth=1, color="b",label ="geglättete Kurve" )
    # ax1 = plt.plot(stauchung,spannung_kn_m2,marker = "s",markevery=15,linewidth=1,color = "red",markersize=5,label = "Original" )

    # plt.plot([],[]," ",label = f"R\u00b2 : {round(r2,3)} " )


    plt.legend(fontsize=16)


    import os
    try:

        data_path = os.getcwd()
        os.mkdir(data_path + "\\Grafik")

    except FileExistsError:
        pass
    plt.savefig(data_path + "\\Grafik" + "\\" + str(probe) + ".png", dpi=500,bbox_inches='tight')  # dpi = 500

    plt.close()

    plt.rcParams["figure.figsize"] = (12, 6)


    fig2 = plt.figure()
    ax2 = fig2.add_subplot(111)

    radius = max(spannung_kn_m2)/2
    y = lambda x : (radius**2 -(x-radius)**2)**0.5
    kreix_x = np.linspace(0, 2 * radius, 500)
    kreis_y = [y(a) for a in np.linspace(0, 2 * radius, 500)]
    ax2.plot(kreix_x,kreis_y,color = "black")
    ax2.set_ylabel("Scherspannung " + r"$\tau$" + " [kN/m²] ",fontsize=20,fontname="Arial")
    ax2.set_xlabel("Einaxiale Druckspannung "+ r"$\sigma$" + " [kN/m²]",fontsize=20,fontname="Arial")
    ax2.plot([radius, radius], np.linspace(0, radius, 2),color="black",linestyle="dashed")
    ax2.plot(np.linspace(0, radius, 2), [radius, radius],color="black",linestyle="dashed")
    ax2.annotate(r"$c_{u}$ = $\tau_{max}$  = " + f"{round(radius,2)} kN/m²".replace(".",",")   ,
        xy=(radius, radius+1), xycoords='data',
        xytext=(0, 0), textcoords='offset points', ha="left", fontsize=20, color="red",fontname="Arial")
    ax2.annotate(r"$q_{u}$ = $\sigma_{max}$ = " + f"{round(radius*2,2)} kN/m²".replace(".",",")   ,
        xy=(2*radius, radius/2), xycoords='data',
        xytext=(0, 0), textcoords='offset points', ha="left", fontsize=20, color="red",fontname="Arial")
    ax2.set_xlim(0,2*radius+10)
    ax2.set_ylim(0,radius+5)
    ax2.xaxis.set_tick_params(labelsize=16)
    ax2.yaxis.set_tick_params(labelsize=16)
    # ax2.plot(stauchung,spannung_kn_m2,marker = "s",markevery=15,linewidth=1,color = "red",markersize=5,label = "Original")
    # ax2.plot(stauchung, new_y, linewidth=1, color="b",label ="geglättete Kurve" )
    # ax2.plot([],[]," ",label = f"R\u00b2 : {round(r2,3)} " )
    # ax2.set_xlabel("Stauchung \u03B5 [%]")
    # ax2.set_ylabel("Druckspannung "+ r"$\sigma$" + " [kN/m²]")

    try:

        data_path = os.getcwd()
        os.mkdir(data_path + "\\Grafik")

    except FileExistsError:
        pass
    plt.savefig(data_path + "\\Grafik" + "\\" + str(probe) + "_2" +".png", dpi=500,bbox_inches='tight')  # dpi = 500

    plt.close()
    versuch_list = read_excel(versuch_list_name, sheet_name="Festigkeiten", header=None).replace(
        nan, "nan")
    festigkeit_proben_nummer = versuch_list[::][1].to_list()
    index_nummer = festigkeit_proben_nummer.index(probe)
    blatt = versuch_list_excel["Festigkeiten"]

    blatt["F"][index_nummer].value = round(radius*2,2)
    blatt["G"][index_nummer].value = round(radius,2)

    versuch_list_excel.save("Neu_"+versuch_list_name)

    return str(round(radius*2,2)).replace(".",","),str(round(return_stauchung,2)).replace(".",",")

def haupt_funktion2(versuch_list,versuch_list_name,layout,block_ort,block_ort_1ax):


        import os
        import shutil
        # dpl_list = dpl_list.to_numpy()
        datei_typen = datei_suchen()  # Die Funktion hier sucht nach Rohdaten.

        i = block_ort  #
        packet = io.BytesIO()
        x = 21 * cm
        y = 29.7 * cm
        seite_größe = (x, y)

        x1 = 11 * cm
        y1 = 23.87 * cm
        kanvas = canvas.Canvas(packet, pagesize=seite_größe)
        zähler = 1


        file_name = versuch_list_name

        data_path = os.getcwd()

        copy_file_name = "Neu_" + versuch_list_name

        kontroll_datei = os.path.join(data_path, copy_file_name)

        if os.path.exists(kontroll_datei):
            versuch_list_excel = openpyxl.load_workbook(copy_file_name)
        else:

            shutil.copy(file_name, copy_file_name)
            versuch_list_excel = openpyxl.load_workbook(copy_file_name)
            print("Neuer Übersicht-Datei erstellt.")


        for j in range(0, versuch_list.shape[0]):


            if versuch_list.to_numpy()[j][block_ort_1ax] == "x":


                #proben_nummer = probennummer(versuch_list, block_ort, index)

                # drehwinkel, spannung = grafik_datei(proben_nummer, datei_typen)
                #
                # try:
                #     tabelle_werten = grafik_darstellung(versuch_list_excel,(drehwinkel, spannung, proben_nummer,x_limit,x_limit_2)
                # except IndexError:
                #     pass
                #
                # H = h_nummerierung_(dpl_list, block_ort)
                # versuch_block = versuchblock(b)[0]  # Es ruft die Funktion Versuchblock auf.
                position = probe_position(versuch_list, j,block_ort)
                # anhang_teil = h_anhanteil_az(dpl_list, block_ort)
                # sondierdatum = sondier_datum(dpl_list, block_ort)

                try:
                    if len(probe_position(versuch_list, j,block_ort)) > 1:
                        radius,stauchung=grafik_darstellung(versuch_list_excel,(versuch_list.to_numpy()[j][block_ort_1ax - 9]))

                        pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                        kanvas.setFont("Arial", 12)
                        # kanvas.setFont("Vera",12)

                        kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")

                        kanvas.drawString(x1, y1, position[0])  # Koordinat von Position
                        kanvas.drawString(9.5 * cm, 23.50 * cm, position[1])
                        kanvas.drawString(6.54 * cm, 23.0 * cm, "Ergebnisse aus dem einaxialen Druckversuch")
                        kanvas.drawString(7.54 * cm, 22.50 * cm, "nach DIN 18136 vom November 2003")

                        kanvas.drawString(17.91 * cm, 25.85 * cm,str(versuch_list.to_numpy()[1][block_ort_1ax-10]))
                        kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                        kanvas.drawString(18.7 * cm, 25.85 * cm, ".6")

                        kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                        kanvas.drawString(16 * cm, 25.49 * cm, "zu Az.: 01/00-" +""+str(versuch_list.to_numpy()[2][block_ort_1ax-10]))
                        kanvas.setFillColorRGB(1, 0, 0)
                        # kanvas.drawString(2.5 * cm, 21.75 * cm, f"Probe :  {proben_nummer}, obere Probenseite ")
                        # kanvas.setFillColorRGB(0, 0, 0)
                        # kanvas.drawString(2.5 * cm, 21.23 * cm,
                        #                   f"Entnahmestelle : {position[0]}, rd. {m} m von der Böschungskante entfernt")
                        # kanvas.drawString(2.5 * cm, 20 * cm, "Sondierdatum : " + f" {sondierdatum}")
                        kanvas.setFillColorRGB(0, 0, 0)

                        kanvas.drawString(2.5 * cm, 21.5 * cm, f"Probe: " + str(versuch_list.to_numpy()[j][block_ort_1ax-9]))
                        kanvas.drawString(2.5 * cm, 21 * cm, f"Entnahmestelle: {position[0]}, rd. " +  f"{str(versuch_list.to_numpy()[j][block_ort_1ax-6])} m von der Böschungskante entfernt".replace(".",","))
                        kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart: REKAL+Asche")
                        kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum: " + str(versuch_list.to_numpy()[0][block_ort_1ax-10].strftime("%d-%m-%Y")))
                        kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum: "+str(versuch_list.to_numpy()[j][block_ort_1ax+1].strftime("%d-%m-%Y")))
                        kanvas.drawString(10.67 * cm, 20.3 * cm, "Anfangshöhe h\u2090: 25,0 cm")
                        kanvas.drawString(10.67 * cm, 19.8 * cm, "Anfangsdurchmesser d\u2090: 11,4 cm")
                        kanvas.drawString(10.67 * cm, 19.3 * cm, "Ausbauwassergehalt: "+str(round(float(versuch_list.to_numpy()[j][block_ort_1ax-3]),2)).replace(".",",")+" %")
                        kanvas.drawString(10.67 * cm, 18.8 * cm, "Vorschubgeschwindigkeit: v = 0,4 mm/min")

                        kanvas.rect(14.7*cm,15.6*cm,5*cm,3*cm)
                        kanvas.setFont("Arial", 11)

                        kanvas.drawString(15.2 * cm, 18 * cm, "Einaxiale Druckfestigkeit ")
                        kanvas.drawString(15.8 * cm, 17.5 * cm, "Bruchstauchung")
                        kanvas.drawString(15.8 * cm, 17 * cm, "q{}".format(get_sub("u")) +  f" = {(radius)} kN/m²")
                        kanvas.drawString(15.8 * cm, 16.5 * cm, "\u03B5{}".format(get_sub("u")) +  f" = {(stauchung)} %")




                        zähler +=1
                        try:
                            import os
                            data_path = os.getcwd()


                            kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_1ax-9])}" + ".png",  2.5* cm, 6.7 * cm,
                                             width=12 * cm, height=12 * cm)

                            kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_1ax-9])}" + "_2"+ ".png",  2.5* cm, 1.1 * cm,
                                             width=12 * cm, height=5.6 * cm)
                        except OSError:
                            pass
                        kanvas.showPage()



                    elif len(probe_position(versuch_list, j,block_ort)) == 1 and probe_position(versuch_list, j,block_ort)[0] != "nan":
                        radius,stauchung=grafik_darstellung(versuch_list_excel,(versuch_list.to_numpy()[j][block_ort_1ax - 9]))
                        pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                        kanvas.setFont("Arial", 12)
                        # kanvas.setFont("Vera",12)

                        kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")

                        kanvas.drawString(x1, y1, position[0])  # Koordinat von Position
                        # kanvas.drawString(8.8 * cm, 23.50 * cm, position[1])
                        kanvas.drawString(6.54 * cm, 23.0 * cm, "Ergebnisse aus dem einaxialen Druckversuch")
                        kanvas.drawString(7.54 * cm, 22.50 * cm, "nach DIN 18136 vom November 2003")

                        kanvas.drawString(17.91 * cm, 25.85 * cm,str(versuch_list.to_numpy()[1][block_ort_1ax-10]))
                        kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                        kanvas.drawString(18.7 * cm, 25.85 * cm, ".6")

                        kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                        kanvas.drawString(16 * cm, 25.49 * cm, "zu Az.: 01/00-" +""+str(versuch_list.to_numpy()[2][block_ort_1ax-10]))
                        kanvas.setFillColorRGB(1, 0, 0)
                        # kanvas.drawString(2.5 * cm, 21.75 * cm, f"Probe :  {proben_nummer}, obere Probenseite ")
                        # kanvas.setFillColorRGB(0, 0, 0)
                        # kanvas.drawString(2.5 * cm, 21.23 * cm,
                        #                   f"Entnahmestelle : {position[0]}, rd. {m} m von der Böschungskante entfernt")
                        # kanvas.drawString(2.5 * cm, 20 * cm, "Sondierdatum : " + f" {sondierdatum}")
                        kanvas.setFillColorRGB(0, 0, 0)


                        kanvas.drawString(2.5 * cm, 21.5 * cm, f"Probe: " + str(versuch_list.to_numpy()[j][block_ort_1ax-9]))
                        kanvas.drawString(2.5 * cm, 21 * cm, f"Entnahmestelle: {position[0]}, rd. " +  f"{str(versuch_list.to_numpy()[j][block_ort_1ax-6])} m von der Böschungskante entfernt".replace(".",","))
                        kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart: REKAL+Asche")
                        kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum: " + str(versuch_list.to_numpy()[0][block_ort_1ax-10].strftime("%d-%m-%Y")))
                        kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum: "+str(versuch_list.to_numpy()[j][block_ort_1ax+1].strftime("%d-%m-%Y")))
                        kanvas.drawString(10.67 * cm, 20.3 * cm, "Anfangshöhe h\u2090: 25,0 cm")
                        kanvas.drawString(10.67 * cm, 19.8 * cm, "Anfangsdurchmesser d\u2090: 11,4 cm")
                        kanvas.drawString(10.67 * cm, 19.3 * cm, "Ausbauwassergehalt: "+str(round(float(versuch_list.to_numpy()[j][block_ort_1ax-3]),2)).replace(".",",")+" %")
                        kanvas.drawString(10.67 * cm, 18.8 * cm, "Vorschubgeschwindigkeit: v = 0,4 mm/min")

                        kanvas.rect(14.7*cm,15.6*cm,5*cm,3*cm)
                        kanvas.setFont("Arial", 11)

                        kanvas.drawString(15.2 * cm, 18 * cm, "Einaxiale Druckfestigkeit ")
                        kanvas.drawString(15.8 * cm, 17.5 * cm, "Bruchstauchung")
                        kanvas.drawString(15.8 * cm, 17 * cm, "q{}".format(get_sub("u")) +  f" = {(radius)} kN/m²")
                        kanvas.drawString(15.8 * cm, 16.5 * cm, "\u03B5{}".format(get_sub("u")) +  f" = {(stauchung)} %")

                        zähler += 1
                        try:
                            import os
                            data_path = os.getcwd()

                            kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_1ax-9])}" + ".png",  2.5* cm, 6.7 * cm,
                                             width=12 * cm, height=12 * cm)

                            kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_1ax-9])}" + "_2"+ ".png",  2.5* cm, 1.1 * cm,
                                             width=12 * cm, height=5.6 * cm)
                        except OSError:
                            pass
                        kanvas.showPage()




                except OSError:
                    pass


            else :

                pass
        kanvas.save()

        packet.seek(0)

        canvas_pdf = PdfFileReader(packet)

        canvas_pdf.getNumPages()

        outfile = PdfFileWriter()

        for z in range(0, canvas_pdf.getNumPages()):
            page = copy.copy(layout.getPage(0))
            page.mergePage(canvas_pdf.getPage(z))

            outfile.addPage(page)

        with open(str(versuch_list.to_numpy()[0][block_ort_1ax-10].strftime("%d-%m-%Y"))+"_1ax_Auswertung.pdf", "wb") as out:
            outfile.write(out)


# plt.show()


# ax1.plot(stauchung, new_y, marker="s", markevery=15)
# plt.plot(stauchung, spannung_kn_m2, marker="s", markevery=15, linewidth=7.0)
# plt.plot(stauchung, new_y, marker="s", markevery=15)


def datei_suchen():
    roh_data_path = os.getcwd() + "\\roh_1ax"

    with os.scandir(roh_data_path) as suchen:
        datei_list = []
        data_typen = []
        for datei_name in suchen:
            datei_list.append(datei_name.name)

        for datei in datei_list:

            if fnmatch.fnmatch(datei, "*.xls") :
                data_typen.append(datei)

    return data_typen



def grafik_datei(proben_nummer, data_typen):
    drehwinkel = {}
    spannungen = {}
    roh_data_path = os.getcwd() + "\\roh_1ax"
    for i in data_typen:

        if i.startswith(f"{proben_nummer}") == True:
            datei = open(roh_data_path + "\\" + f"{i}", "r")

            roh_data = datei.readlines()
            grafik_datei = []
            grafik_datei_2 = []

            for j in range(len(roh_data)):
                grafik_datei.append((roh_data[j][3:10].replace(",", ".")))

                grafik_datei_2.append((roh_data[j][15:21].replace(",", ".")))
            drehwinkel[f"{i}"] = grafik_datei[11:]
            spannungen[f"{i}"] = grafik_datei_2[11:]

            datei.close()

    return drehwinkel, spannungen

def probe_position(versuch_list, j,block_ort):

    if block_ort == 0 :
        position = str(versuch_list[3][j]).split(",")

        if len(position) == 1 and position[0] == "nan":

            return None



        elif len(position) == 1 and position[0] != "nan":

            return (position)


        else:
            return "Böschung " + str(position[0]), f"({position[1]})"

    elif block_ort == 1 :
        position = str(versuch_list[23][j]).split(",")

        if len(position) == 1 and position[0] == "nan":

            return None



        elif len(position) == 1 and position[0] != "nan":

            return (position)


        else:
            return "Böschung " + str(position[0]), f"({position[1]})"


def h_nummerierung_(dpl_list,block_ort):


    if block_ort == 1 :

        return dpl_list[6][1]

    elif block_ort == 2 :

        return  dpl_list[35][1]





def h_anhanteil_az(dpl_list,block_ort):

    if block_ort == 1:

        return dpl_list[5][1]

    elif block_ort == 2:

        return dpl_list[34][1]


def versuchblock(b):

    if b == 0 :
        return ".1" , "Plateau"

    elif b  == 1 :
        return ".2","Böschung Süd"

    elif b == 2 :
        return ".3" , "Böschung Nord"

def sondier_datum(dpl_list,block_ort):

    if block_ort == 1:

        return dpl_list[11][1].strftime("%d-%m-%Y")

    elif block_ort == 2:

        return dpl_list[31][1].strftime("%d-%m-%Y")

def entnahmedatum(versuch_list, i):
    a = 20 * i

    if type(versuch_list.columns[a]) == type(" "):
        print("Datum ist gleich wie vorherige , bitte entnahme datum ändern")
    return (versuch_list.columns[a]).strftime("%d-%m-%Y")


def get_sub(x):
    normal = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789+-=()"
    sub_s = "ₐ₈CDₑբGₕᵢⱼₖₗₘₙₒₚQᵣₛₜᵤᵥwₓᵧZₐ♭꜀ᑯₑբ₉ₕᵢⱼₖₗₘₙₒₚ૧ᵣₛₜᵤᵥwₓᵧ₂₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎"
    res = x.maketrans(''.join(normal), ''.join(sub_s))
    return x.translate(res)


# durchmesser_cm = 11.40
# anfangshöhe_cm = 25.00
# anfangsvolumen_cm3 = round(np.pi*(durchmesser_cm*0.5)**2*anfangshöhe_cm,2)
# weg_intern = daten[9]
# weg_ab_kontakt = weg_intern-weg_intern[0]
# kraft_kN = daten[1]
# massgebliche_querschnitt_m2 = anfangsvolumen_cm3/(anfangshöhe_cm-weg_ab_kontakt/10)/10000
# spannung_kn_m2 = kraft_kN/massgebliche_querschnitt_m2
#
# stauchung = ((weg_ab_kontakt/10)/anfangshöhe_cm)*100
#
#
# curve = np.polyfit(stauchung,spannung_kn_m2,10)
# poly = np.poly1d(curve)
# print(curve,poly)
#
# new_y = [poly(i) for i in stauchung]
# # plt.plot(stauchung,new_y)
# plt.plot(stauchung,spannung_kn_m2,marker = "s",markevery=15,linewidth=7.0)
# plt.plot(stauchung,new_y,marker = "s",markevery=15)
#
#
# plt.show()




print("""

Das Programm bietet die Auswertungen der Testergebnisse von 1ax zu berechnen.


!!!Hinweis: Wenn das Programm eine Fehlermeldung ausgibt, muss es neu gestartet werden.!!!

!!!Testergebnisse müssen in einem stabilen Excel-Format vorliegen. !!!


!!!Änderungen sollten in der beigefügten Excel-Datei vorgenommen werden. Neue Testergebnisse sollten auf derselben Datei vorbereitet werden.!!!

1 - ) Dateiname von Versuchlist eingeben

2- ) Geben bitte Auswahl ein !!!

     Alle Auswertungen von einer Liste zu erstellen : 1 eingeben 

3 - ) Drücken Sie 0, um das Programm zu schließen

""")
try:







    versuch_list_name = input("Bitte Exceldatei von Versucliste eingeben !!! ")
    versuch_list = read_excel(versuch_list_name, sheet_name="Uebersicht",header = None).replace(
        nan, "nan")

    versuch_list[11] = to_datetime(versuch_list[11], errors='coerce', dayfirst=True)
    versuch_list[13] = to_datetime(versuch_list[13], errors='coerce', dayfirst=True)
    versuch_list[14] = to_datetime(versuch_list[14], errors='coerce', dayfirst=True)

    versuch_list[18] = to_datetime(versuch_list[18], errors='coerce', dayfirst=True)
    versuch_list[31] = to_datetime(versuch_list[31], errors='coerce', dayfirst=True)
    versuch_list[33] = to_datetime(versuch_list[33], errors='coerce', dayfirst=True)
    versuch_list[34] = to_datetime(versuch_list[34], errors='coerce', dayfirst=True)

    versuch_list[14] = versuch_list[14].replace(nan, "-")
    versuch_list[13] = versuch_list[13].replace(nan, "-")
    versuch_list[33] = versuch_list[33].replace(nan, "-")
    versuch_list[34] = versuch_list[34].replace(nan, "-")
except FileNotFoundError:
    print("Excel-Datei nicht gefunden, überprüft mal bitte den Dateinamen."


          "Bitte starten das Programm neu"

          )

try:

    layout_path = os.getcwd()

    layout = PdfFileReader(f"{layout_path}" + "\\" + "Layout_1ax.pdf", "rb")


except FileNotFoundError:
    print("Layout-Datei nicht gefunden, überprüft mal bitte den Dateinamen."
          "Bitte starten das Programm neu"

          ""
          "")



while True:
    auswahl = int(input(" Auswahl :  "))

    if type(auswahl) == type(" ") or auswahl > 1 or auswahl < 0:

        print("Es wurde ein ungültiger Wert eingegeben, bitte versuchen Sie es erneut.")

    else:



        if auswahl == 1:
            try:
                block_ort = int(input("Bitte Blocknummer von Probennummer eingeben 1. oder 2. Block")) - 1

                if block_ort == 0:
                    block_ort_1ax = 10
                elif block_ort == 1:
                    block_ort_1ax = 30

                haupt_funktion2(versuch_list, versuch_list_name, layout, block_ort, block_ort_1ax)
                print("""   

                *************
                Prozess erfolgreich abgeschlossen !!!
                *************

                """
                      )

                print(""" * * Drücken Sie 0, um das Programm zu schließen! 

                          * * einen neuen PDF zu erstellen 1 drücken 


                      """)
            except (PermissionError,NameError):

                print("Bitte Excel-Datei schlisessen oder Dateiname kontrolieren")



        elif auswahl == 2 :

            try:

                x_limit = int(input("Bitte X-Achse Anfang-Limit eingeben (Vorgegebene Wert : 1850)"))
                x_limit_2 = int(input("Bitte X-Achse End-Limit eingeben (Vorgegebene Wert : 1950)"))
                y_limit = int(input("Bitte Y-Achse Anfang-Limit eingeben (Vorgegebene Wert : 0)"))
                y_limit2 = int(input("Bitte Y-Achse End-Limit eingeben (Vorgegebene Wert : 100)"))
            except NameError:

                print("Falsche Eingabe")

        elif auswahl == 0:

            print("Skript erfolgreich abgeschlossen !!!"
                  ""
                  "Tschüss")
            break






