import copy
import copy
import os
import matplotlib.pyplot as plt
import numpy as np
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch, cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PyPDF2 import PdfFileWriter, PdfFileReader
from reportlab.platypus import Frame, Image
from pandas import read_excel
from numpy import nan
import io
import fnmatch
from PyPDF2 import PdfFileWriter, PdfFileReader
from pandas import read_excel
from pandas import to_datetime
from pandas import NaT
from matplotlib.patches import Rectangle
from numpy import nan
# from reportlab.graphics.shapes import Rect,Line
# from matplotlib.transforms import Affine2D
# import matplotlib.ticker as plticker
import openpyxl
import shutil


def grafik_darstellung(versuch_list_excel,versuch_list_name,probe,proben_seite):

    if proben_seite == 1 :

        proben_seite_loc = "-O"
    elif proben_seite == 2 :

        proben_seite_loc = "-U"


    import os

    versuch_list = read_excel(versuch_list_name, sheet_name="Festigkeiten", header=None).replace(
        nan, "nan")
    roh_data_path = os.getcwd() + "\\roh_pen"
    ergebnisse_1 = read_excel(roh_data_path+ "/" +str(probe)+proben_seite_loc +".xlsm", sheet_name="Ergebnisse_1",header=None).round(decimals=6)
    ergebnisse_2 = read_excel(roh_data_path+ "/" +str(probe)+proben_seite_loc +".xlsm", sheet_name="Ergebnisse_2",header=None).round(decimals=6)
    ergebnisse_3 = read_excel(roh_data_path+ "/" +str(probe)+proben_seite_loc +".xlsm", sheet_name="Ergebnisse_3",header=None).round(decimals=6)

    kraft_kn_1 = ergebnisse_1[1][2::].to_numpy()
    kraft_kn_2 = ergebnisse_2[1][2::].to_numpy()
    kraft_kn_3 = ergebnisse_3[1][2::].to_numpy()

    max_y = [max(kraft_kn_1),max(kraft_kn_2),max(kraft_kn_3)]
    max_y = max(max_y)

    penetrationstiefe_mm_1 =ergebnisse_1[2][2::].to_numpy()
    penetrationstiefe_mm_2 =ergebnisse_2[2][2::].to_numpy()
    penetrationstiefe_mm_3 =ergebnisse_3[2][2::].to_numpy()

    arbeit_1 = np.zeros(len(penetrationstiefe_mm_1))
    arbeit_2 = np.zeros(len(penetrationstiefe_mm_2))
    arbeit_3 = np.zeros(len(penetrationstiefe_mm_3))

    for i in range(0,len(penetrationstiefe_mm_1)-1) :

        if ((penetrationstiefe_mm_1[i+1] - penetrationstiefe_mm_1[i])* (kraft_kn_1[i+1] + kraft_kn_1[i])/2) < 0 :

            arbeit_1[i+1] = 0

        else :


            arbeit_1[i+1] = round(((penetrationstiefe_mm_1[i+1] - penetrationstiefe_mm_1[i])* (kraft_kn_1[i+1] + kraft_kn_1[i])/2),4)

    for i in range(0,len(penetrationstiefe_mm_2)-1) :

        if ((penetrationstiefe_mm_2[i + 1] - penetrationstiefe_mm_2[i]) * (kraft_kn_2[i + 1] + kraft_kn_2[i]) / 2) < 0:

            arbeit_2[i + 1] = 0

        else:

            arbeit_2[i + 1] = round(((penetrationstiefe_mm_2[i + 1] - penetrationstiefe_mm_2[i]) * (kraft_kn_2[i + 1] + kraft_kn_2[i]) / 2),4)

    for i in range(0, len(penetrationstiefe_mm_3) - 1):

        if ((penetrationstiefe_mm_3[i + 1] - penetrationstiefe_mm_3[i]) * (kraft_kn_3[i + 1] + kraft_kn_3[i]) / 2) < 0:

            arbeit_3[i + 1] = 0

        else:

            arbeit_3[i + 1] = round(((penetrationstiefe_mm_3[i + 1] - penetrationstiefe_mm_3[i]) * (kraft_kn_3[i + 1] + kraft_kn_3[i]) / 2),4)

    arbeit_1_sum = sum(arbeit_1)
    arbeit_2_sum = sum(arbeit_2)
    arbeit_3_sum = sum(arbeit_3)

    scalar_faktor = 272.84

    arbeit_1 = round(scalar_faktor*arbeit_1_sum,3)
    arbeit_2 = round(scalar_faktor*arbeit_2_sum,3)
    arbeit_3 = round(scalar_faktor*arbeit_3_sum,3)

    festigkeit_proben_nummer = versuch_list[::][1].to_list()
    index_nummer = festigkeit_proben_nummer.index(probe)
    blatt = versuch_list_excel["Festigkeiten"]

    if proben_seite == 1 :

        blatt["P"][index_nummer].value = arbeit_1
        blatt["Q"][index_nummer].value = arbeit_2
        blatt["R"][index_nummer].value = arbeit_3
    elif proben_seite == 2 :
        blatt["S"][index_nummer].value = arbeit_1
        blatt["T"][index_nummer].value = arbeit_2
        blatt["U"][index_nummer].value = arbeit_3



    versuch_list_excel.save("Neu_"+versuch_list_name)

    plt.rcParams["figure.figsize"] = (14, 13)
    plt.rc('xtick', labelsize=16)
    plt.rc('ytick', labelsize=16)
    fig = plt.figure()
    ax1 = fig.add_subplot(111)
    # plt.rc('xtick', labelsize=12)
    # plt.rc('ytick', labelsize=12)

    #
    ax1.plot(penetrationstiefe_mm_1[2::],kraft_kn_1[2::],marker = "s",markevery=15,linewidth=1.5,color = "k",markersize=10,label = "Messung 1: "+f"{arbeit_1} kN/m²".replace(".",","))
    ax1.plot(penetrationstiefe_mm_2[2::],kraft_kn_2[2::],marker = "x",markevery=15,linewidth=1.5,color = "r",markersize=10,label = "Messung 2: "+f"{arbeit_2} kN/m²".replace(".",","))
    ax1.plot(penetrationstiefe_mm_3[2::],kraft_kn_3[2::],marker = "d",markevery=15,linewidth=1.5,color = "b",markersize=10,label = "Messung 3: "+f"{arbeit_3} kN/m²".replace(".",","))
    ax1.add_patch(Rectangle((28,0.8*(max_y*3)),55,0.5*(max_y*3), color="black", fill=False))


    ax1.set_xlim(0,55)
    ax1.set_ylim(0,max_y*3)
    ax1.annotate("Spitze vollständig im Probenmaterial ",
                 xy=(28.4,0.82*(max_y*3)), xycoords='data',
                 xytext=(0, 0), textcoords='offset points', ha="left", fontsize=15, color="black", fontname="Arial")
    ax1.annotate(r"$w_{50}$ = $\frac{1}{\rm A[mm^2] * s_{50} [mm]}$" + r'$\int_{\rm s = 0 [mm]}^{\rm s_{50} = 50 [mm]}$' + r"F$_{\rm c}$ [kN] ds 1)",
                 xy=(28.4,0.9*(max_y*3)), xycoords='data',
                 xytext=(0, 0), textcoords='offset points', ha="left", fontsize=18, color="black", fontname="Arial")
    plt.figtext(0.5, 0.03, "1) Der $w_{50}$-Wert bezeichnet die für die Penetration erforderliche Arbeit (Kraft [F] x Weg [L]) \n bezogen auf das durch den Penetrator verdrängte Volumen [L³].", wrap=True, horizontalalignment='center', fontsize=16)
    # ax1.add_patch(Rectangle((0,0) ,20     , 20, fill=False, hatch=h))
    ax1.set_xlabel("s [mm]",fontsize = 18)
    ax1.set_ylabel(r"F$_{\rm c}$ [kN]",fontsize = 18) ### Um nicht kursiv zu schreiben.

    ax1.grid()
    # ax1.xaxis.set_tick_params(labelsize=16)
    # ax1.yaxis.set_tick_params(labelsize=16)
    #
    #
    plt.legend(fontsize=16,loc = 2)
    #
    #
    import os
    try:

        data_path = os.getcwd()
        os.mkdir(data_path + "\\Grafik")

    except FileExistsError:
        pass
    # plt.savefig(data_path + "\\Grafik" + "\\" + str(probe) + ".svg", dpi=500,bbox_inches='tight',format="svg")  # dpi = 500
    plt.savefig(data_path + "\\Grafik" + "\\" + str(probe) +f"{proben_seite}"+".png", dpi=500, bbox_inches='tight',
                format="png")  # dpi = 500
    plt.close()


def haupt_funktion2(versuch_list,versuch_list_name,layout,block_ort,block_ort_pen):

        import os
        # dpl_list = dpl_list.to_numpy()
        datei_typen = datei_suchen()  # Die Funktion hier sucht nach Rohdaten.

        i = block_ort  #
        packet = io.BytesIO()
        x = 21 * cm
        y = 29.7 * cm
        seite_größe = (x, y)

        x1 = 11 * cm
        y1 = 23.87 * cm
        kanvas = canvas.Canvas(packet, pagesize=seite_größe)
        zähler = 1
        nat = NaT


        file_name = versuch_list_name

        data_path = os.getcwd()

        copy_file_name = "Neu_"+versuch_list_name

        kontroll_datei = os.path.join(data_path,copy_file_name)

        if os.path.exists(kontroll_datei):
            versuch_list_excel = openpyxl.load_workbook(copy_file_name)
        else :

            shutil.copy(file_name,copy_file_name)
            versuch_list_excel = openpyxl.load_workbook(copy_file_name)
            print("Neuer Übersicht-Datei erstellt.")


        for j in range(0, versuch_list.shape[0]):


            if versuch_list.to_numpy()[j][block_ort_pen] == "x":



                if ((versuch_list.to_numpy()[j][block_ort_pen+1] != "-" )  and (versuch_list.to_numpy()[j][block_ort_pen+2] != "-"))  ==  True   :

                    position = probe_position(versuch_list, j, block_ort)
                    for i in range(1,3) :

                            

                            if i == 1:
                                nummer_versuchblock = nummerierung_versuchblock(versuch_list, block_ort, int(versuch_list.to_numpy()[j][block_ort_pen - 11]))
                                probe_name = str(versuch_list.to_numpy()[j][block_ort_pen-11]) + ", obere Probenseite"
                                grafik_darstellung(versuch_list_excel,versuch_list_name,versuch_list.to_numpy()[j][block_ort_pen-11],i)
                                grafik_name = str(versuch_list.to_numpy()[j][block_ort_pen-11]) + "-O"

                            elif i == 2 :

                                probe_name = str(versuch_list.to_numpy()[j][block_ort_pen - 11]) + ", untere Probenseite"
                                nummer_versuchblock = nummerierung_versuchblock(versuch_list, block_ort, int(versuch_list.to_numpy()[j][block_ort_pen - 11]))+1


                                grafik_darstellung(versuch_list_excel,versuch_list_name,versuch_list.to_numpy()[j][block_ort_pen-11], i)
                                grafik_name = str(versuch_list.to_numpy()[j][block_ort_pen - 11]) + "-U"

                            try:
                                if len(probe_position(versuch_list, j,block_ort)) > 1:
                                    pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                                    kanvas.setFont("Arial", 12)
                                    # kanvas.setFont("Vera",12)

                                    kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")

                                    kanvas.drawString(x1, y1,position[0])  # Koordinat von Position
                                    kanvas.drawString(9.5 * cm, 23.50 * cm, position[1])
                                    kanvas.drawString(6.54 * cm, 23.0 * cm, "       Ergebnisse des Penetrationsversuchs")
                                    kanvas.drawString(17.80 * cm, 25.85 * cm,str(versuch_list.to_numpy()[1][block_ort_pen-12]))
                                    kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                                    kanvas.drawString(18.58 * cm, 25.85 * cm, ".5")
                                    kanvas.drawString(18.98 * cm, 25.85 * cm, versuchblock(position))
                                    kanvas.drawString(19.38 * cm, 25.85 * cm, ".0"+str(nummer_versuchblock))
                                    # kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                                    kanvas.drawString(16 * cm, 25.49 * cm, "zu Az.: 01/00-" +""+str(versuch_list.to_numpy()[2][block_ort_pen-12]))
                                    kanvas.setFillColorRGB(1, 0, 0)


                                    kanvas.drawString(2.5 * cm, 21.5 * cm, f"Probe: " + probe_name)
                                    kanvas.setFillColorRGB(0, 0, 0)

                                    kanvas.drawString(2.5 * cm, 21 * cm, f"Entnahmestelle:         {position[0]}, rd. " +  f"{str(versuch_list.to_numpy()[j][block_ort_pen-8])} m".replace(".",","))
                                    kanvas.drawString(2.5*cm,20.5*cm,"entfernt von der Böschungskante.")
                                    kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart:                    REKAL+Asche")
                                    kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum:        " + str(versuch_list.to_numpy()[0][block_ort_pen-12].strftime("%d-%m-%Y")))
                                    kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum:            "+str(versuch_list.to_numpy()[j][block_ort_pen+i].strftime("%d-%m-%Y")))
                                    kanvas.drawString(2.5 * cm, 17.8 * cm, "Wassergehalt:             "+str(round(float(versuch_list.to_numpy()[j][block_ort_pen-5]),2)).replace(".",",")+" %")
                                    kanvas.drawString(2.5 * cm, 17.3 * cm, "Feuchtdichte:              "+str(round(float(versuch_list.to_numpy()[j][block_ort_pen-4]),2)).replace(".",",")+" g/cm³")
                                    kanvas.drawString(2.5 * cm, 16.8 * cm, "Trockendichte:            "+str(round(float(versuch_list.to_numpy()[j][block_ort_pen-3]),2)).replace(".",",")+" g/cm³")
                                    # kanvas.setFillColorRGB(1, 0, 0)
                                    # kanvas.drawString(13 * cm, 20.4 * cm, "Penetratorspitze")
                                    # kanvas.setFillColorRGB(0, 0, 0)
                                    #
                                    # kanvas.drawString(13 * cm, 19.9 * cm, "Vorschub: v = 5 mm/min")
                                    # kanvas.drawString(13 * cm, 19.4 * cm, "Spitzenwinkel: 90°")
                                    # kanvas.drawString(13 * cm, 18.9 * cm, "Spitzenquerschnitt: A ≈ 0,79 cm²")




                                    #
                                    zähler +=1
                                    try:
                                        import os
                                        data_path = os.getcwd()


                                        kanvas.drawImage(data_path + "\\Grafik\\" + f"{grafik_name}" + ".png",  2.5* cm, 2.2*cm,
                                                         width=16 * cm, height=14 * cm)

                                    #     kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_pen-11])}" + "_2"+ ".png",  2.5* cm, 1.1 * cm,
                                    #                      width=14 * cm, height=5.6 * cm)
                                    except OSError:
                                        pass
                                    kanvas.showPage()



                                elif len(probe_position(versuch_list, j,block_ort)) == 1 and probe_position(versuch_list, j,block_ort)[0] != "nan":

                                    pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                                    kanvas.setFont("Arial", 12)
                                    # kanvas.setFont("Vera",12)
                                    kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")
                                    kanvas.drawString(x1, y1, position[0])  # Koordinat von Position
                                    # kanvas.drawString(8.8 * cm, 23.50 * cm, position[1])
                                    kanvas.drawString(6.54 * cm, 23.0 * cm, "     Ergebnisse des Penetrationsversuchs")

                                    kanvas.drawString(17.80 * cm, 25.85 * cm, str(versuch_list.to_numpy()[1][block_ort_pen - 12]))
                                    kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                                    kanvas.drawString(18.58 * cm, 25.85 * cm, ".5")
                                    kanvas.drawString(18.98 * cm, 25.85 * cm, versuchblock(position))

                                    kanvas.drawString(19.38 * cm, 25.85 * cm, ".0"+str(nummer_versuchblock))
                                    # kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                                    kanvas.drawString(16 * cm, 25.49 * cm,"zu Az.: 01/00-" + "" + str(versuch_list.to_numpy()[2][block_ort_pen - 12]))
                                    kanvas.setFillColorRGB(1, 0, 0)
                                    # kanvas.drawString(2.5 * cm, 21.75 * cm, f"Probe :  {proben_nummer}, obere Probenseite ")
                                    # kanvas.setFillColorRGB(0, 0, 0)
                                    # kanvas.drawString(2.5 * cm, 21.23 * cm,
                                    #                   f"Entnahmestelle : {position[0]}, rd. {m} m von der Böschungskante entfernt")
                                    # kanvas.drawString(2.5 * cm, 20 * cm, "Sondierdatum : " + f" {sondierdatum}")
                                    kanvas.setFillColorRGB(1, 0, 0)


                                    kanvas.drawString(2.5 * cm, 21.5 * cm, f"Probe: " + probe_name)
                                    kanvas.setFillColorRGB(0, 0, 0)

                                    kanvas.drawString(2.5 * cm, 21 * cm, f"Entnahmestelle:         {position[0]}, rd. " +  f"{str(versuch_list.to_numpy()[j][block_ort_pen-8])} m".replace(".",","))
                                    kanvas.drawString(2.5*cm,20.5*cm,"entfernt von der Böschungskante.")
                                    kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart:                    REKAL+Asche")
                                    kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum:        " + str(versuch_list.to_numpy()[0][block_ort_pen-12].strftime("%d-%m-%Y")))
                                    kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum:            "+str(versuch_list.to_numpy()[j][block_ort_pen+i].strftime("%d-%m-%Y")))
                                    kanvas.drawString(2.5 * cm, 17.8 * cm, "Wassergehalt:             "+str(round(float(versuch_list.to_numpy()[j][block_ort_pen-5]),2)).replace(".",",")+" %")
                                    kanvas.drawString(2.5 * cm, 17.3 * cm, "Feuchtdichte:              "+str(round(float(versuch_list.to_numpy()[j][block_ort_pen-4]),2)).replace(".",",")+" g/cm³")
                                    kanvas.drawString(2.5 * cm, 16.8 * cm, "Trockendichte:            "+str(round(float(versuch_list.to_numpy()[j][block_ort_pen-3]),2)).replace(".",",")+" g/cm³")


                                    #
                                    # kanvas.setFillColorRGB(1, 0, 0)
                                    # kanvas.drawString(13 * cm, 20.4 * cm, "Penetratorspitze")
                                    # kanvas.setFillColorRGB(0, 0, 0)
                                    #
                                    # kanvas.drawString(13 * cm, 19.9 * cm, "Vorschub: v = 5 mm/min")
                                    # kanvas.drawString(13 * cm, 19.4 * cm, "Spitzenwinkel: 90°")
                                    # kanvas.drawString(13 * cm, 18.9 * cm, "Spitzenquerschnitt: A ≈ 0,79 cm²")




                                    zähler += 1
                                    try:
                                        import os
                                        data_path = os.getcwd()

                                        kanvas.drawImage(data_path + "\\Grafik\\" + f"{grafik_name}" + ".png",  2.5* cm, 2.2*cm,
                                                         width=16 * cm, height=14 * cm)

                                    except OSError:
                                        pass
                                    kanvas.showPage()




                            except OSError:
                                pass
                elif  ((versuch_list.to_numpy()[j][block_ort_pen+1]  != "-")  and (versuch_list.to_numpy()[j][block_ort_pen+2] == "-")) == True  :


                    position = probe_position(versuch_list, j, block_ort)
                    grafik_darstellung(versuch_list_excel,versuch_list_name,versuch_list.to_numpy()[j][block_ort_pen - 11], 1)

                    try:
                        if len(probe_position(versuch_list, j, block_ort)) > 1:

                            pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                            kanvas.setFont("Arial", 12)
                            # kanvas.setFont("Vera",12)
                            nummer_versuchblock = nummerierung_versuchblock(versuch_list, block_ort, int(versuch_list.to_numpy()[j][block_ort_pen - 11]))
                            kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")
                            kanvas.drawString(x1, y1, position[0])  # Koordinat von Position
                            kanvas.drawString(9.5 * cm, 23.50 * cm, position[1])
                            kanvas.drawString(6.54 * cm, 23.0 * cm, "       Ergebnisse des Penetrationsversuchs")
                            kanvas.drawString(17.80 * cm, 25.85 * cm,str(versuch_list.to_numpy()[1][block_ort_pen - 12]))
                            kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                            kanvas.drawString(18.58 * cm, 25.85 * cm, ".5")
                            kanvas.drawString(18.98 * cm, 25.85 * cm, versuchblock(position))
                            kanvas.drawString(19.38 * cm, 25.85 * cm, ".0"+str(nummer_versuchblock))
                            # kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                            kanvas.drawString(16 * cm, 25.49 * cm, "zu Az.: 01/00-" + "" + str(
                                versuch_list.to_numpy()[2][block_ort_pen - 12]))
                            kanvas.setFillColorRGB(1, 0, 0)

                            kanvas.drawString(2.5 * cm, 21.5 * cm,
                                              f"Probe: " + str(versuch_list.to_numpy()[j][block_ort_pen - 11]) + ", obere Probenseite")
                            kanvas.setFillColorRGB(0, 0, 0)

                            kanvas.drawString(2.5 * cm, 21 * cm,
                                              f"Entnahmestelle:         {position[0]}, rd. " + f"{str(versuch_list.to_numpy()[j][block_ort_pen - 8])} m".replace(
                                                  ".", ","))
                            kanvas.drawString(2.5 * cm, 20.5 * cm, "entfernt von der Böschungskante.")
                            kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart:                    REKAL+Asche")
                            kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum:        " + str(
                                versuch_list.to_numpy()[0][block_ort_pen - 12].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum:            "+str(versuch_list.to_numpy()[j][block_ort_pen+1].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 17.8 * cm, "Wassergehalt:             " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 5]), 2)).replace(".",
                                                                                                        ",") + " %")
                            kanvas.drawString(2.5 * cm, 17.3 * cm, "Feuchtdichte:              " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 4]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")
                            kanvas.drawString(2.5 * cm, 16.8 * cm, "Trockendichte:            " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 3]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")

                            # kanvas.setFillColorRGB(1, 0, 0)
                            # kanvas.drawString(13 * cm, 20.4 * cm, "Penetratorspitze")
                            # kanvas.setFillColorRGB(0, 0, 0)
                            #
                            # kanvas.drawString(13 * cm, 19.9 * cm, "Vorschub: v = 5 mm/min")
                            # kanvas.drawString(13 * cm, 19.4 * cm, "Spitzenwinkel: 90°")
                            # kanvas.drawString(13 * cm, 18.9 * cm, "Spitzenquerschnitt: A ≈ 0,79 cm²")

                            #
                            zähler += 1

                            try:
                                import os
                                data_path = os.getcwd()

                                
                                kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_pen-11])}" +"-O" + ".png",  2.5* cm, 2.2*cm,
                                                 width=16 * cm, height=14 * cm)

                            except OSError:
                                pass
                            kanvas.showPage()



                        elif len(probe_position(versuch_list, j, block_ort)) == 1 and probe_position(versuch_list, j, block_ort)[0] != "nan":
                            nummer_versuchblock = nummerierung_versuchblock(versuch_list, block_ort, int(versuch_list.to_numpy()[j][block_ort_pen - 11]))
                            pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                            kanvas.setFont("Arial", 12)
                            # kanvas.setFont("Vera",12)

                            kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")

                            kanvas.drawString(x1, y1, position[0])  # Koordinat von Position
                            # kanvas.drawString(8.8 * cm, 23.50 * cm, position[1])
                            kanvas.drawString(6.54 * cm, 23.0 * cm, "     Ergebnisse des Penetrationsversuchs")

                            kanvas.drawString(17.80* cm, 25.85 * cm,str(versuch_list.to_numpy()[1][block_ort_pen - 12]))
                            kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                            kanvas.drawString(18.58 * cm, 25.85 * cm, ".5")
                            kanvas.drawString(18.98 * cm, 25.85 * cm, versuchblock(position))
                            kanvas.drawString(19.38 * cm, 25.85 * cm, ".0"+str(nummer_versuchblock))
                            # kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                            kanvas.drawString(16 * cm, 25.49 * cm, "zu Az.: 01/00-" + "" + str(
                                versuch_list.to_numpy()[2][block_ort_pen - 12]))
                            kanvas.setFillColorRGB(1, 0, 0)
                            # kanvas.drawString(2.5 * cm, 21.75 * cm, f"Probe :  {proben_nummer}, obere Probenseite ")
                            # kanvas.setFillColorRGB(0, 0, 0)
                            # kanvas.drawString(2.5 * cm, 21.23 * cm,
                            #                   f"Entnahmestelle : {position[0]}, rd. {m} m von der Böschungskante entfernt")
                            # kanvas.drawString(2.5 * cm, 20 * cm, "Sondierdatum : " + f" {sondierdatum}")
                            kanvas.setFillColorRGB(1, 0, 0)

                            kanvas.drawString(2.5 * cm, 21.5 * cm,
                                              f"Probe: " + str(versuch_list.to_numpy()[j][block_ort_pen - 11]))
                            kanvas.setFillColorRGB(0, 0, 0)

                            kanvas.drawString(2.5 * cm, 21 * cm,
                                              f"Entnahmestelle:         {position[0]}, rd. " + f"{str(versuch_list.to_numpy()[j][block_ort_pen - 8])} m".replace(
                                                  ".", ","))
                            kanvas.drawString(2.5 * cm, 20.5 * cm, "entfernt von der Böschungskante.")
                            kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart:                    REKAL+Asche")
                            kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum:        " + str(
                                versuch_list.to_numpy()[0][block_ort_pen - 12].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum:            "+str(versuch_list.to_numpy()[j][block_ort_pen+1].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 17.8 * cm, "Wassergehalt:             " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 5]), 2)).replace(".",
                                                                                                        ",") + " %")
                            kanvas.drawString(2.5 * cm, 17.3 * cm, "Feuchtdichte:              " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 4]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")
                            kanvas.drawString(2.5 * cm, 16.8 * cm, "Trockendichte:            " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 3]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")

                            # kanvas.setFillColorRGB(1, 0, 0)
                            # kanvas.drawString(13 * cm, 20.4 * cm, "Penetratorspitze")
                            # kanvas.setFillColorRGB(0, 0, 0)
                            #
                            # kanvas.drawString(13 * cm, 19.9 * cm, "Vorschub: v = 5 mm/min")
                            # kanvas.drawString(13 * cm, 19.4 * cm, "Spitzenwinkel: 90°")
                            # kanvas.drawString(13 * cm, 18.9 * cm, "Spitzenquerschnitt: A ≈ 0,79 cm²")

                            zähler += 1
                            try:
                                import os
                                data_path = os.getcwd()

                                kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_pen-11])}"+"-O" + ".png",  2.5* cm, 2.2*cm,
                                                 width=16 * cm, height=14 * cm)

                            except OSError:
                                pass
                            kanvas.showPage()





                    except OSError:
                        pass

                elif (versuch_list.to_numpy()[j][block_ort_pen+1]  == "-")  and (versuch_list.to_numpy()[j][block_ort_pen+2] != "-") == True :

                    position = probe_position(versuch_list, j, block_ort)
                    grafik_darstellung(versuch_list_excel,versuch_list_name,versuch_list.to_numpy()[j][block_ort_pen - 11], 2)
                    for i in range(1,3):
                        if (type(versuch_list.to_numpy()[j][block_ort_pen+i] ) != type(nat)) :

                            datum = versuch_list.to_numpy()[j][block_ort_pen+i]


                    try:
                        if len(probe_position(versuch_list, j, block_ort)) > 1:
                            nummer_versuchblock = nummerierung_versuchblock(versuch_list, block_ort, int(versuch_list.to_numpy()[j][block_ort_pen - 11]))
                            pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                            kanvas.setFont("Arial", 12)
                            # kanvas.setFont("Vera",12)

                            kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")

                            kanvas.drawString(x1, y1, position[0])  # Koordinat von Position
                            kanvas.drawString(9.5 * cm, 23.50 * cm, position[1])
                            kanvas.drawString(6.54 * cm, 23.0 * cm, "       Ergebnisse des Penetrationsversuchs")
                            kanvas.drawString(17.80 * cm, 25.85 * cm,str(versuch_list.to_numpy()[1][block_ort_pen - 12]))
                            kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                            kanvas.drawString(18.58 * cm, 25.85 * cm, ".5")
                            kanvas.drawString(18.98 * cm, 25.85 * cm, versuchblock(position))
                            kanvas.drawString(19.38 * cm, 25.85 * cm, ".0"+str(nummer_versuchblock))
                            # kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                            kanvas.drawString(16 * cm, 25.49 * cm, "zu Az.: 01/00-" + "" + str(
                                versuch_list.to_numpy()[2][block_ort_pen - 12]))
                            kanvas.setFillColorRGB(1, 0, 0)

                            kanvas.drawString(2.5 * cm, 21.5 * cm,
                                              f"Probe: " + str(versuch_list.to_numpy()[j][block_ort_pen - 11]) + ", untere Probenseite")
                            kanvas.setFillColorRGB(0, 0, 0)

                            kanvas.drawString(2.5 * cm, 21 * cm,
                                              f"Entnahmestelle:         {position[0]}, rd. " + f"{str(versuch_list.to_numpy()[j][block_ort_pen - 8])} m".replace(
                                                  ".", ","))
                            kanvas.drawString(2.5 * cm, 20.5 * cm, "entfernt von der Böschungskante.")
                            kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart:                    REKAL+Asche")
                            kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum:        " + str(
                                versuch_list.to_numpy()[0][block_ort_pen - 12].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum:            "+str(versuch_list.to_numpy()[j][block_ort_pen+2].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 17.8 * cm, "Wassergehalt:             " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 5]), 2)).replace(".",
                                                                                                        ",") + " %")
                            kanvas.drawString(2.5 * cm, 17.3 * cm, "Feuchtdichte:              " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 4]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")
                            kanvas.drawString(2.5 * cm, 16.8 * cm, "Trockendichte:            " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 3]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")

                            # kanvas.setFillColorRGB(1, 0, 0)
                            # kanvas.drawString(13 * cm, 20.4 * cm, "Penetratorspitze")
                            # kanvas.setFillColorRGB(0, 0, 0)
                            #
                            # kanvas.drawString(13 * cm, 19.9 * cm, "Vorschub: v = 5 mm/min")
                            # kanvas.drawString(13 * cm, 19.4 * cm, "Spitzenwinkel: 90°")
                            # kanvas.drawString(13 * cm, 18.9 * cm, "Spitzenquerschnitt: A ≈ 0,79 cm²")

                            #
                            zähler += 1

                            try:
                                import os
                                data_path = os.getcwd()


                                kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_pen-11])}" +"-U" + ".png",  2.5* cm, 2.2*cm,
                                                 width=16 * cm, height=14 * cm)


                            except OSError:
                                pass
                            kanvas.showPage()



                        elif len(probe_position(versuch_list, j, block_ort)) == 1 and \
                                probe_position(versuch_list, j, block_ort)[0] != "nan":
                            nummer_versuchblock = nummerierung_versuchblock(versuch_list, block_ort, int(versuch_list.to_numpy()[j][block_ort_pen - 11]))
                            pdfmetrics.registerFont(TTFont('Arial', 'Fonts/arial.ttf'))
                            kanvas.setFont("Arial", 12)
                            # kanvas.setFont("Vera",12)

                            kanvas.drawString(8.2 * cm, y1, "REKAL Halde,")

                            kanvas.drawString(x1, y1, position[0])  # Koordinat von Position
                            # kanvas.drawString(8.8 * cm, 23.50 * cm, position[1])
                            kanvas.drawString(6.54 * cm, 23.0 * cm, "     Ergebnisse des Penetrationsversuchs")

                            kanvas.drawString(17.80 * cm, 25.85 * cm,str(versuch_list.to_numpy()[1][block_ort_pen - 12]))
                            kanvas.drawString(16 * cm, 25.85 * cm, "Anhang:")
                            kanvas.drawString(18.58 * cm, 25.85 * cm, ".5")
                            kanvas.drawString(18.98 * cm, 25.85 * cm, versuchblock(position))
                            kanvas.drawString(19.38 * cm, 25.85 * cm, ".0"+str(nummer_versuchblock))
                            # kanvas.drawString(19.1 * cm, 25.85 * cm, "."+str(zähler))

                            kanvas.drawString(16 * cm, 25.49 * cm, "zu Az.: 01/00-" + "" + str(
                                versuch_list.to_numpy()[2][block_ort_pen - 12]))
                            kanvas.setFillColorRGB(1, 0, 0)
                            # kanvas.drawString(2.5 * cm, 21.75 * cm, f"Probe :  {proben_nummer}, obere Probenseite ")
                            # kanvas.setFillColorRGB(0, 0, 0)
                            # kanvas.drawString(2.5 * cm, 21.23 * cm,
                            #                   f"Entnahmestelle : {position[0]}, rd. {m} m von der Böschungskante entfernt")
                            # kanvas.drawString(2.5 * cm, 20 * cm, "Sondierdatum : " + f" {sondierdatum}")
                            kanvas.setFillColorRGB(1, 0, 0)

                            kanvas.drawString(2.5 * cm, 21.5 * cm,
                                              f"Probe: " + str(versuch_list.to_numpy()[j][block_ort_pen - 11]))
                            kanvas.setFillColorRGB(0, 0, 0)

                            kanvas.drawString(2.5 * cm, 21 * cm,
                                              f"Entnahmestelle:         {position[0]}, rd. " + f"{str(versuch_list.to_numpy()[j][block_ort_pen - 8])} m".replace(
                                                  ".", ","))
                            kanvas.drawString(2.5 * cm, 20.5 * cm, "entfernt von der Böschungskante.")
                            kanvas.drawString(2.5 * cm, 19.8 * cm, "Bodenart:                    REKAL+Asche")
                            kanvas.drawString(2.5 * cm, 19.3 * cm, "Entnahmedatum:        " + str(
                                versuch_list.to_numpy()[0][block_ort_pen - 12].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 18.8 * cm, "Sondierdatum:            "+str(versuch_list.to_numpy()[j][block_ort_pen+2].strftime("%d-%m-%Y")))
                            kanvas.drawString(2.5 * cm, 17.8 * cm, "Wassergehalt:             " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 5]), 2)).replace(".",
                                                                                                        ",") + " %")
                            kanvas.drawString(2.5 * cm, 17.3 * cm, "Feuchtdichte:              " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 4]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")
                            kanvas.drawString(2.5 * cm, 16.8 * cm, "Trockendichte:            " + str(
                                round(float(versuch_list.to_numpy()[j][block_ort_pen - 3]), 2)).replace(".",
                                                                                                        ",") + " g/cm³")

                            # kanvas.setFillColorRGB(1, 0, 0)
                            # kanvas.drawString(13 * cm, 20.4 * cm, "Penetratorspitze")
                            # kanvas.setFillColorRGB(0, 0, 0)
                            #
                            # kanvas.drawString(13 * cm, 19.9 * cm, "Vorschub: v = 5 mm/min")
                            # kanvas.drawString(13 * cm, 19.4 * cm, "Spitzenwinkel: 90°")
                            # kanvas.drawString(13 * cm, 18.9 * cm, "Spitzenquerschnitt: A ≈ 0,79 cm²")

                            zähler += 1
                            try:
                                import os
                                data_path = os.getcwd()

                                kanvas.drawImage(data_path + "\\Grafik\\" + f"{str(versuch_list.to_numpy()[j][block_ort_pen-11])}" +"-U"  + ".png",  2.5* cm, 2.2*cm,width=16 * cm, height=14 * cm)


                            except OSError:
                                pass
                            kanvas.showPage()





                    except OSError:
                        pass

            else :

                pass
        kanvas.save()

        packet.seek(0)

        canvas_pdf = PdfFileReader(packet)

        canvas_pdf.getNumPages()

        outfile = PdfFileWriter()

        for z in range(0, canvas_pdf.getNumPages()):
            page = copy.copy(layout.getPage(0))
            page.mergePage(canvas_pdf.getPage(z))

            outfile.addPage(page)
        try :
            with open(str(versuch_list.to_numpy()[0][block_ort_pen-12].strftime("%d-%m-%Y"))+"_PEN_Auswertung.pdf", "wb") as out:
                outfile.write(out)
        except PermissionError :
            with open(str(versuch_list.to_numpy()[0][block_ort_pen-12].strftime("%d-%m-%Y"))+"_PEN_Auswertung.pdf"+"_2", "wb") as out:
                outfile.write(out)
# plt.show()


# ax1.plot(stauchung, new_y, marker="s", markevery=15)
# plt.plot(stauchung, spannung_kn_m2, marker="s", markevery=15, linewidth=7.0)
# plt.plot(stauchung, new_y, marker="s", markevery=15)


def datei_suchen():
    roh_data_path = os.getcwd() + "\\roh_pen"

    with os.scandir(roh_data_path) as suchen:
        datei_list = []
        data_typen = []
        for datei_name in suchen:
            datei_list.append(datei_name.name)

        for datei in datei_list:

            if fnmatch.fnmatch(datei, "*.xlsm") :
                data_typen.append(datei)

    return data_typen



def grafik_datei(proben_nummer, data_typen):
    drehwinkel = {}
    spannungen = {}
    roh_data_path = os.getcwd() + "\\roh_pen"
    for i in data_typen:

        if i.startswith(f"{proben_nummer}") == True:
            datei = open(roh_data_path + "\\" + f"{i}", "r")

            roh_data = datei.readlines()
            grafik_datei = []
            grafik_datei_2 = []

            for j in range(len(roh_data)):
                grafik_datei.append((roh_data[j][3:10].replace(",", ".")))

                grafik_datei_2.append((roh_data[j][15:21].replace(",", ".")))
            drehwinkel[f"{i}"] = grafik_datei[11:]
            spannungen[f"{i}"] = grafik_datei_2[11:]

            datei.close()

    return drehwinkel, spannungen

def probe_position(versuch_list, j,block_ort):

    if block_ort == 0 :
        position = str(versuch_list[3][j]).split(",")

        if len(position) == 1 and position[0] == "nan":

            return None



        elif len(position) == 1 and position[0] != "nan":

            return (position)


        else:
            return "Böschung " + str(position[0]), f"({position[1]})"

    elif block_ort == 1 :
        position = str(versuch_list[23][j]).split(",")

        if len(position) == 1 and position[0] == "nan":

            return None



        elif len(position) == 1 and position[0] != "nan":

            return (position)


        else:
            return "Böschung " + str(position[0]), f"({position[1]})"


def h_nummerierung_(dpl_list,block_ort):


    if block_ort == 1 :

        return dpl_list[6][1]

    elif block_ort == 2 :

        return  dpl_list[35][1]





def h_anhanteil_az(dpl_list,block_ort):

    if block_ort == 1:

        return dpl_list[5][1]

    elif block_ort == 2:

        return dpl_list[34][1]


def versuchblock(b):

    if b == 0 :
        return ".1" , "Plateau"

    elif b  == 1 :
        return ".2","Böschung Süd"

    elif b == 2 :
        return ".3" , "Böschung Nord"

def sondier_datum(dpl_list,block_ort):

    if block_ort == 1:

        return dpl_list[11][1].strftime("%d-%m-%Y")

    elif block_ort == 2:

        return dpl_list[31][1].strftime("%d-%m-%Y")

def entnahmedatum(versuch_list, i):
    a = 20 * i

    if type(versuch_list.columns[a]) == type(" "):
        print("Datum ist gleich wie vorherige , bitte entnahme datum ändern")
    return (versuch_list.columns[a]).strftime("%d-%m-%Y")


def get_sub(x):
    normal = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789+-=()"
    sub_s = "ₐ₈CDₑբGₕᵢⱼₖₗₘₙₒₚQᵣₛₜᵤᵥwₓᵧZₐ♭꜀ᑯₑբ₉ₕᵢⱼₖₗₘₙₒₚ૧ᵣₛₜᵤᵥwₓᵧ₂₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎"
    res = x.maketrans(''.join(normal), ''.join(sub_s))
    return x.translate(res)

def nummerierung_versuchblock(versuch_list, block_ort, proben_nummer):
    if block_ort == 0:
        block = 3

    elif block_ort == 1:
        block = 23

    versuch_list = versuch_list.to_numpy()
    platau_versuchblock = []
    sho_versuchblock = []
    shu_versuchblock = []
    nho_versuchblock = []
    nhu_versuchblock = []

    for i in range(versuch_list.shape[0]):

        if ((versuch_list[i][block] == 'Plateau') and (versuch_list[i][block+9] == "x" )) ==  True:

            if ((versuch_list[i][block+10] != "-" )  and (versuch_list[i][block+11] != "-"))  ==  True :
                platau_versuchblock.append(versuch_list[i][block - 2])
                platau_versuchblock.append(versuch_list[i][block - 2])

            # elif ((versuch_list[i][block+9]  != "-")  and (versuch_list[i][block+9] == "-")) == True  :
            else :
                platau_versuchblock.append(versuch_list[i][block - 2])


            platau_versuchblock = [x for x in platau_versuchblock if x != 'nan']

        elif ((versuch_list[i][block] == 'Süd,50 m über Fuß') and (versuch_list[i][block+9] == "x" )) == True:


            if ((versuch_list[i][block+10]  != "-")  and (versuch_list[i][block+11] != "-")) == True  :
                sho_versuchblock.append(versuch_list[i][block - 2])
                sho_versuchblock.append(versuch_list[i][block - 2])
            else :
                sho_versuchblock.append(versuch_list[i][block - 2])

            sho_versuchblock = [x for x in sho_versuchblock if x != 'nan']


        elif (versuch_list[i][block] == 'Süd,direkt am Fuß' ) and (versuch_list[i][block+9] == "x" ):

            if ((versuch_list[i][block + 10] != "-") and (versuch_list[i][block + 11] != "-")) == True:
                shu_versuchblock.append(versuch_list[i][block - 2])
                shu_versuchblock.append(versuch_list[i][block - 2])
            else :
                shu_versuchblock.append(versuch_list[i][block - 2])
            shu_versuchblock = [x for x in shu_versuchblock if x != 'nan']

        elif (versuch_list[i][block] == 'Nord,50 m über Fuß') and (versuch_list[i][block+9] == "x" ):
            if ((versuch_list[i][block + 10] != "-") and (versuch_list[i][block + 11] != "-")) == True:

                nho_versuchblock.append(versuch_list[i][block - 2])
                nho_versuchblock.append(versuch_list[i][block - 2])
            else :
                nho_versuchblock.append(versuch_list[i][block - 2])
            nho_versuchblock = [x for x in nho_versuchblock if x != 'nan']
        elif (versuch_list[i][block] == 'Nord,direkt am Fuß') and (versuch_list[i][block+9] == "x" ):

            if ((versuch_list[i][block + 10] != "-") and (versuch_list[i][block + 11] != "-")) == True:

                nhu_versuchblock.append(versuch_list[i][block - 2])
                nhu_versuchblock.append(versuch_list[i][block - 2])
            else :
                nhu_versuchblock.append(versuch_list[i][block - 2])
            nhu_versuchblock = [x for x in nhu_versuchblock if x != 'nan']


    for x in (platau_versuchblock, shu_versuchblock, nho_versuchblock, sho_versuchblock,
              nhu_versuchblock):
        for y in x:

            if y == proben_nummer:
                return int(x.index(y) + 1)

def versuchblock(position):
    if position[0] == "Plateau":
        return ".1"

    elif position[0] + position[1] == "Böschung Süd(50 m über Fuß)":
        return ".2"

    elif position[0] + position[1] == "Böschung Süd(direkt am Fuß)":
        return ".3"
    elif position[0] + position[1] == "Böschung Nord(50 m über Fuß)":
        return ".4"

    elif position[0] + position[1] == 'Böschung Nord(direkt am Fuß)':
        return ".5"


print("""

Das Programm bietet die Auswertungen der Testergebnisse von Penatrationversuch zu berechnen.


!!!Hinweis: Wenn das Programm eine Fehlermeldung ausgibt, muss es neu gestartet werden.!!!

!!!Testergebnisse müssen in einem stabilen Excel-Format vorliegen. !!!


!!!Änderungen sollten in der beigefügten Excel-Datei vorgenommen werden. Neue Testergebnisse sollten auf derselben Datei vorbereitet werden.!!!

1 - ) Dateiname von Versuchlist eingeben

2- ) Geben bitte Auswahl ein !!!

     Alle Auswertungen von einer Liste zu erstellen : 1 eingeben 

3 - ) Drücken Sie 0, um das Programm zu schließen

""")
try:







    versuch_list_name = input("Bitte Exceldatei von Versucliste eingeben !!! ")
    versuch_list = read_excel(versuch_list_name, sheet_name="Uebersicht",header = None).replace(
        nan, "nan")

    versuch_list[11] = to_datetime(versuch_list[11], errors='coerce', dayfirst=True)
    versuch_list[13] = to_datetime(versuch_list[13], errors='coerce', dayfirst=True)
    versuch_list[14] = to_datetime(versuch_list[14], errors='coerce', dayfirst=True)

    versuch_list[18] = to_datetime(versuch_list[18], errors='coerce', dayfirst=True)
    versuch_list[31] = to_datetime(versuch_list[31], errors='coerce', dayfirst=True)
    versuch_list[33] = to_datetime(versuch_list[33], errors='coerce', dayfirst=True)
    versuch_list[34] = to_datetime(versuch_list[34], errors='coerce', dayfirst=True)

    versuch_list[14] = versuch_list[14].replace(nan, "-")
    versuch_list[13] = versuch_list[13].replace(nan, "-")
    versuch_list[33] = versuch_list[33].replace(nan, "-")
    versuch_list[34] = versuch_list[34].replace(nan, "-")
except FileNotFoundError:
    print("Excel-Datei nicht gefunden, überprüft mal bitte den Dateinamen."


          "Bitte starten das Programm neu"

          )

try:

    layout_path = os.getcwd()

    layout = PdfFileReader(f"{layout_path}" + "\\" + "Layout_PEN.pdf", "rb")


except FileNotFoundError:
    print("Layout-Datei nicht gefunden, überprüft mal bitte den Dateinamen."
          "Bitte starten das Programm neu"

          ""
          "")



while True:
    auswahl = int(input(" Auswahl :  "))

    if type(auswahl) == type(" ") or auswahl > 1 or auswahl < 0:

        print("Es wurde ein ungültiger Wert eingegeben, bitte versuchen Sie es erneut.")

    else:



        if auswahl == 1:
            try:
                block_ort = int(input("Bitte Blocknummer von Probennummer eingeben 1. oder 2. Block")) - 1

                if block_ort == 0:
                    block_ort_pen = 12
                elif block_ort == 1:
                    block_ort_pen = 32

                haupt_funktion2(versuch_list,versuch_list_name,layout,block_ort,block_ort_pen)

                print("""   

                *************
                Prozess erfolgreich abgeschlossen !!!
                *************

                """
                      )

                print(""" * * Drücken Sie 0, um das Programm zu schließen! 

                          * * einen neuen PDF zu erstellen 1 drücken 


                      """)
            except (PermissionError,NameError):

                print("Bitte Excel-Datei schlisessen oder Dateiname kontrolieren")



        elif auswahl == 2 :

            try:

                x_limit = int(input("Bitte X-Achse Anfang-Limit eingeben (Vorgegebene Wert : 1850)"))
                x_limit_2 = int(input("Bitte X-Achse End-Limit eingeben (Vorgegebene Wert : 1950)"))
                y_limit = int(input("Bitte Y-Achse Anfang-Limit eingeben (Vorgegebene Wert : 0)"))
                y_limit2 = int(input("Bitte Y-Achse End-Limit eingeben (Vorgegebene Wert : 100)"))
            except NameError:

                print("Falsche Eingabe")

        elif auswahl == 0:

            print("Skript erfolgreich abgeschlossen !!!"
                  ""
                  "Tschüss")
            break


# a = datei_suchen()
# b = [x for x in a if x.endswith("-O.xlsm") or x.endswith("-U.xlsm") ]
# pen_proben={}
# for i in range(len(b) - 1):
#     if b[i][0:5] == b[i + 1][0:5]:
#         pen_proben[b[i][0:5]] = [b[i], b[i + 1]
#                                  ]



